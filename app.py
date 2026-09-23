"""Client-ready Streamlit app for video/audio transcription by Team Fahad."""

from __future__ import annotations

import io
import os
import re
import shutil
import tempfile
import time
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import streamlit as st
from docx import Document
from google import genai
from google.genai import errors
from moviepy import AudioFileClip, VideoFileClip


APP_TITLE = "Video/Audio Transcriber by Team Fahad"
TRANSCRIBE_MODEL = "gemini-3.5-transcribe"
CONTENT_MODEL = "gemini-3.8-flash"
MAX_UPLOAD_MB = 500
MAX_ANALYSIS_CHARS = 120_000

VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv"}
AUDIO_MIME_TYPES = {
    ".mp3": "audio/mpeg",
    ".wav": "audio/wav",
    ".m4a": "audio/m4a",
    ".aac": "audio/aac",
    ".ogg": "audio/ogg",
    ".flac": "audio/flac",
}

LANGUAGES = {
    "Auto detect / mixed language": [],
    "Bangla (Bangladesh)": ["bn-BD"],
    "English (US)": ["en-US"],
    "English (UK)": ["en-GB"],
    "Arabic": ["ar-EG"],
    "Hindi": ["hi-IN"],
}

TRANSLATION_LANGUAGES = [
    "Bangla",
    "English",
    "Arabic",
    "Hindi",
]


@dataclass
class WordInfo:
    text: str
    speaker: str | None
    start: float | None
    end: float | None


@dataclass
class Segment:
    start: float
    end: float
    text: str
    speaker: str | None = None


def secret_or_env(name: str) -> str | None:
    """Read a secret from Streamlit first, then from environment variables."""
    try:
        value = st.secrets.get(name)
        if value:
            return str(value)
    except Exception:
        pass
    value = os.getenv(name)
    return value if value else None


def safe_name(name: str) -> str:
    stem = Path(name).stem
    stem = re.sub(r"[^\w\-]+", "_", stem, flags=re.UNICODE).strip("_")
    return stem[:80] or "transcript"


def parse_offset(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if text.endswith("s"):
        text = text[:-1]
    try:
        return float(text)
    except ValueError:
        return None


def format_clock(seconds: float) -> str:
    seconds = max(0.0, seconds)
    total = int(seconds)
    h = total // 3600
    m = (total % 3600) // 60
    s = total % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


def format_srt_time(seconds: float) -> str:
    millis = max(0, round(seconds * 1000))
    h, remainder = divmod(millis, 3_600_000)
    m, remainder = divmod(remainder, 60_000)
    s, ms = divmod(remainder, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def format_vtt_time(seconds: float) -> str:
    return format_srt_time(seconds).replace(",", ".")


def smart_join(tokens: list[str]) -> str:
    """Join token-like words while avoiding spaces before punctuation."""
    text = ""
    for token in tokens:
        token = token.strip()
        if not token:
            continue
        if not text:
            text = token
        elif re.match(r"^[,.;:!?%\)\]\}।]", token):
            text += token
        elif text.endswith(("(", "[", "{", "“", '"', "'")):
            text += token
        else:
            text += " " + token
    return text.strip()


def wait_for_file(client: genai.Client, uploaded_file: Any, timeout_seconds: int = 180) -> Any:
    """Wait for an uploaded Gemini file to become usable when state is exposed."""
    state_obj = getattr(uploaded_file, "state", None)
    if state_obj is None:
        return uploaded_file

    deadline = time.monotonic() + timeout_seconds
    current = uploaded_file
    while True:
        state = getattr(getattr(current, "state", None), "name", str(getattr(current, "state", "")))
        state = state.upper()
        if state in {"ACTIVE", "STATE_ACTIVE", ""}:
            return current
        if "FAILED" in state:
            raise RuntimeError("Gemini could not process the uploaded media file.")
        if time.monotonic() >= deadline:
            raise TimeoutError("Gemini did not finish preparing the uploaded file in time.")
        time.sleep(2)
        current = client.files.get(name=current.name)


def extract_word_annotations(interaction: Any) -> list[WordInfo]:
    words: list[WordInfo] = []
    for step in getattr(interaction, "steps", []) or []:
        for content in getattr(step, "content", []) or []:
            for annotation in getattr(content, "annotations", []) or []:
                if getattr(annotation, "type", None) != "word_info":
                    continue
                words.append(
                    WordInfo(
                        text=str(getattr(annotation, "text", "") or "").strip(),
                        speaker=(str(getattr(annotation, "speaker", "") or "").strip() or None),
                        start=parse_offset(getattr(annotation, "start_offset", None)),
                        end=parse_offset(getattr(annotation, "end_offset", None)),
                    )
                )
    return [w for w in words if w.text]


def make_segments(words: list[WordInfo], max_duration: float = 7.0, max_chars: int = 88) -> list[Segment]:
    """Turn word annotations into readable subtitle-sized segments."""
    timed = [w for w in words if w.start is not None and w.end is not None]
    if not timed:
        return []

    segments: list[Segment] = []
    current: list[WordInfo] = []

    def flush() -> None:
        nonlocal current
        if not current:
            return
        segments.append(
            Segment(
                start=current[0].start or 0.0,
                end=current[-1].end or current[0].start or 0.0,
                text=smart_join([w.text for w in current]),
                speaker=current[0].speaker,
            )
        )
        current = []

    for word in timed:
        if not current:
            current.append(word)
            continue

        current_start = current[0].start or 0.0
        duration = (word.end or current_start) - current_start
        speaker_changed = bool(word.speaker and current[0].speaker and word.speaker != current[0].speaker)
        prospective = smart_join([w.text for w in current] + [word.text])
        sentence_end = bool(re.search(r"[.!?।][\"'”’)]?$", current[-1].text))

        if speaker_changed or duration > max_duration or len(prospective) > max_chars:
            flush()
        elif sentence_end and duration >= 2.0:
            flush()

        current.append(word)

    flush()
    return [s for s in segments if s.text]


def speaker_transcript(segments: list[Segment]) -> str:
    if not segments:
        return ""
    lines = []
    for segment in segments:
        speaker = segment.speaker or "Speaker"
        speaker = speaker.replace("spk_", "Speaker ").replace("SPK_", "Speaker ")
        lines.append(f"[{format_clock(segment.start)}] {speaker}: {segment.text}")
    return "\n".join(lines)


def segments_to_srt(segments: list[Segment]) -> str:
    blocks = []
    for idx, segment in enumerate(segments, start=1):
        speaker = ""
        if segment.speaker:
            label = segment.speaker.replace("spk_", "Speaker ").replace("SPK_", "Speaker ")
            speaker = f"{label}: "
        blocks.append(
            f"{idx}\n{format_srt_time(segment.start)} --> {format_srt_time(segment.end)}\n"
            f"{speaker}{segment.text}\n"
        )
    return "\n".join(blocks)


def segments_to_vtt(segments: list[Segment]) -> str:
    blocks = ["WEBVTT\n"]
    for segment in segments:
        speaker = ""
        if segment.speaker:
            label = segment.speaker.replace("spk_", "Speaker ").replace("SPK_", "Speaker ")
            speaker = f"{label}: "
        blocks.append(
            f"{format_vtt_time(segment.start)} --> {format_vtt_time(segment.end)}\n"
            f"{speaker}{segment.text}\n"
        )
    return "\n".join(blocks)


def media_duration(path: Path, is_video: bool) -> float | None:
    try:
        if is_video:
            with VideoFileClip(str(path)) as clip:
                return float(clip.duration or 0.0)
        with AudioFileClip(str(path)) as clip:
            return float(clip.duration or 0.0)
    except Exception:
        return None


def prepare_audio(source: Path, extension: str, folder: Path) -> tuple[Path, str, float | None]:
    """Return (audio_path, mime_type, duration_seconds)."""
    if extension in VIDEO_EXTENSIONS:
        duration = media_duration(source, is_video=True)
        audio_path = folder / "extracted_audio.mp3"
        try:
            with VideoFileClip(str(source)) as video:
                if video.audio is None:
                    raise ValueError("This video does not contain an audio track.")
                video.audio.write_audiofile(str(audio_path), codec="libmp3lame", logger=None)
        except ValueError:
            raise
        except Exception as exc:
            raise RuntimeError(f"Could not extract audio from this video: {exc}") from exc
        return audio_path, "audio/mpeg", duration

    duration = media_duration(source, is_video=False)
    return source, AUDIO_MIME_TYPES[extension], duration


def transcribe_media(
    uploaded: Any,
    api_key: str,
    language_codes: list[str],
    mode: str,
    custom_vocabulary: list[str],
) -> dict[str, Any]:
    extension = Path(uploaded.name).suffix.lower()
    if extension not in VIDEO_EXTENSIONS and extension not in AUDIO_MIME_TYPES:
        raise ValueError("Unsupported file format.")

    upload_size_mb = getattr(uploaded, "size", 0) / (1024 * 1024)
    if upload_size_mb > MAX_UPLOAD_MB:
        raise ValueError(f"File is too large. Maximum allowed size is {MAX_UPLOAD_MB} MB.")

    client = genai.Client(api_key=api_key)
    remote_file = None

    try:
        with tempfile.TemporaryDirectory(prefix="team_fahad_transcriber_") as temp_dir:
            folder = Path(temp_dir)
            source = folder / f"source{extension}"
            with source.open("wb") as destination:
                uploaded.seek(0)
                shutil.copyfileobj(uploaded, destination)

            if source.stat().st_size == 0:
                raise ValueError("The uploaded file is empty.")

            audio_path, mime_type, duration = prepare_audio(source, extension, folder)

            detailed = mode == "Detailed subtitles + speakers"
            duration_limit = 30 * 60 if detailed else 60 * 60
            if duration and duration > duration_limit + 1:
                minutes = duration_limit // 60
                raise ValueError(
                    f"This mode supports up to about {minutes} minutes per file. "
                    "Please split the media into smaller parts and try again."
                )

            remote_file = client.files.upload(file=str(audio_path), config={"mime_type": mime_type})
            remote_file = wait_for_file(client, remote_file)

            transcription_config: dict[str, Any] = {"language_codes": language_codes}
            if mode == "Smart clean transcript":
                transcription_config["mode"] = "smart"
            elif mode == "Detailed subtitles + speakers":
                transcription_config["mode"] = {
                    "type": "verbatim",
                    "diarization_mode": "speaker",
                    "timestamp_granularities": ["word"],
                }
            else:
                transcription_config["mode"] = {"type": "verbatim"}

            if custom_vocabulary and mode != "Detailed subtitles + speakers":
                transcription_config["custom_vocabulary"] = custom_vocabulary[:100]

            input_item = {
                "type": "audio",
                "uri": remote_file.uri,
                "mime_type": getattr(remote_file, "mime_type", mime_type) or mime_type,
            }

            last_error: Exception | None = None
            interaction = None
            for attempt in range(2):
                try:
                    interaction = client.interactions.create(
                        model=TRANSCRIBE_MODEL,
                        input=[input_item],
                        generation_config={"transcription_config": transcription_config},
                    )
                    break
                except errors.APIError as exc:
                    last_error = exc
                    code = getattr(exc, "code", None)
                    if code != 503 or attempt == 1:
                        raise
                    time.sleep(5)

            if interaction is None:
                raise RuntimeError(f"Transcription request failed: {last_error}")

            transcript = str(getattr(interaction, "output_text", "") or "").strip()
            if not transcript:
                raise RuntimeError("Gemini returned an empty transcription.")

            words = extract_word_annotations(interaction)
            segments = make_segments(words)

            return {
                "transcript": transcript,
                "words": words,
                "segments": segments,
                "speaker_transcript": speaker_transcript(segments),
                "srt": segments_to_srt(segments),
                "vtt": segments_to_vtt(segments),
                "duration": duration,
                "mode": mode,
                "language": next((name for name, codes in LANGUAGES.items() if codes == language_codes), "Auto detect"),
            }
    finally:
        if remote_file is not None:
            try:
                client.files.delete(name=remote_file.name)
            except Exception:
                pass
        try:
            client.close()
        except Exception:
            pass


def ai_text(api_key: str, prompt: str) -> str:
    client = genai.Client(api_key=api_key)
    try:
        result = client.interactions.create(model=CONTENT_MODEL, input=prompt)
        text = str(getattr(result, "output_text", "") or "").strip()
        if not text:
            raise RuntimeError("Gemini returned an empty result.")
        return text
    finally:
        try:
            client.close()
        except Exception:
            pass


def transcript_for_analysis(result: dict[str, Any]) -> str:
    detailed = result.get("speaker_transcript") or ""
    base = detailed if detailed else result.get("transcript", "")
    if len(base) > MAX_ANALYSIS_CHARS:
        base = base[:MAX_ANALYSIS_CHARS] + "\n\n[Transcript truncated for this AI add-on.]"
    return base


def generate_summary(api_key: str, result: dict[str, Any]) -> str:
    transcript = transcript_for_analysis(result)
    return ai_text(
        api_key,
        """Summarize the following transcript for a client. Keep the summary accurate and useful.
Use the same main language as the transcript. Include:
1) a short executive summary,
2) 5-10 key points,
3) action items or decisions only if they are actually present.
Do not invent facts.\n\nTRANSCRIPT:\n""" + transcript,
    )


def generate_translation(api_key: str, result: dict[str, Any], target: str) -> str:
    transcript = transcript_for_analysis(result)
    return ai_text(
        api_key,
        f"""Translate the transcript below into {target}. Preserve meaning, names, numbers, and paragraph structure.
If timestamp/speaker labels are present, preserve them. Do not summarize.\n\nTRANSCRIPT:\n{transcript}""",
    )


def generate_content_pack(api_key: str, result: dict[str, Any]) -> str:
    transcript = transcript_for_analysis(result)
    timestamp_note = (
        "The transcript includes timestamps, so create accurate YouTube chapter suggestions from them."
        if result.get("speaker_transcript")
        else "The transcript has no reliable timestamps, so do not invent chapter times."
    )
    return ai_text(
        api_key,
        f"""Turn this transcript into a practical creator/client content pack. Use the transcript's main language unless English is clearly better for a field.
{timestamp_note}
Return these sections:
- 10 strong title ideas
- Short description
- Full SEO-friendly description
- Key takeaways
- 5 short hooks for reels/shorts
- Suggested social caption
- Relevant tags/keywords/hashtags
- YouTube chapters only when reliable timestamps are present
- One clear CTA
Do not invent claims that are not in the transcript.\n\nTRANSCRIPT:\n{transcript}""",
    )


def make_docx_bytes(title: str, sections: list[tuple[str, str]]) -> bytes:
    doc = Document()
    doc.add_heading(title, level=0)
    for heading, body in sections:
        if not body:
            continue
        doc.add_heading(heading, level=1)
        for paragraph in body.split("\n"):
            doc.add_paragraph(paragraph)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_zip_bytes(base_name: str, result: dict[str, Any], extras: dict[str, str]) -> bytes:
    memory = io.BytesIO()
    with zipfile.ZipFile(memory, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"{base_name}_transcript.txt", result.get("transcript", ""))
        if result.get("speaker_transcript"):
            zf.writestr(f"{base_name}_speaker_transcript.txt", result["speaker_transcript"])
        if result.get("srt"):
            zf.writestr(f"{base_name}.srt", result["srt"])
        if result.get("vtt"):
            zf.writestr(f"{base_name}.vtt", result["vtt"])
        for key, value in extras.items():
            if value:
                zf.writestr(f"{base_name}_{key}.txt", value)
        docx_sections = [("Transcript", result.get("transcript", ""))]
        if result.get("speaker_transcript"):
            docx_sections.append(("Speaker Transcript", result["speaker_transcript"]))
        docx_sections.extend((key.replace("_", " ").title(), value) for key, value in extras.items() if value)
        zf.writestr(f"{base_name}_complete.docx", make_docx_bytes(APP_TITLE, docx_sections))
    return memory.getvalue()


def reset_outputs() -> None:
    for key in ["result", "summary", "translation", "content_pack", "translation_target", "source_name"]:
        st.session_state.pop(key, None)


def render_access_gate() -> bool:
    access_code = secret_or_env("APP_ACCESS_CODE")
    if not access_code:
        return True
    if st.session_state.get("access_granted"):
        return True
    st.title(APP_TITLE)
    st.caption("Private client preview")
    entered = st.text_input("Access code", type="password")
    if st.button("Open app", type="primary"):
        if entered == access_code:
            st.session_state.access_granted = True
            st.rerun()
        else:
            st.error("Incorrect access code.")
    return False


st.set_page_config(page_title=APP_TITLE, page_icon="🎙️", layout="wide")

THEMES = {
    "Midnight Neon": {
        "bg": "#07111f", "panel": "#0d1728", "panel2": "#101d31", "text": "#f7fbff",
        "muted": "#93a4bc", "accent": "#7c5cff", "accent2": "#00d7ff", "border": "rgba(148,163,184,.18)",
        "hero1": "#161d46", "hero2": "#0d2944", "glow": "rgba(124,92,255,.30)",
    },
    "Ocean Glass": {
        "bg": "#06151d", "panel": "#0a202b", "panel2": "#0d2835", "text": "#effcff",
        "muted": "#99bcc7", "accent": "#11b5e4", "accent2": "#35f0c1", "border": "rgba(125,211,252,.18)",
        "hero1": "#0b2d3c", "hero2": "#0a3b47", "glow": "rgba(17,181,228,.26)",
    },
    "Emerald Studio": {
        "bg": "#071611", "panel": "#0b2019", "panel2": "#0f2a21", "text": "#f3fff9",
        "muted": "#9ab9aa", "accent": "#28d17c", "accent2": "#8df3b9", "border": "rgba(134,239,172,.18)",
        "hero1": "#123526", "hero2": "#0c2c28", "glow": "rgba(40,209,124,.26)",
    },
    "Retro Wave": {
        "bg": "#120817", "panel": "#1a0d25", "panel2": "#241135", "text": "#fff4df",
        "muted": "#d2a8c9", "accent": "#ff4fd8", "accent2": "#00e6ff", "border": "rgba(255,118,194,.18)",
        "hero1": "#3b165d", "hero2": "#241357", "glow": "rgba(255,79,216,.28)",
    },
}

# Theme is intentionally stored in session state so switching it reruns the app instantly.
if "ui_theme" not in st.session_state:
    st.session_state.ui_theme = "Midnight Neon"

with st.sidebar:
    st.markdown("<div class='side-kicker'>APPEARANCE</div>", unsafe_allow_html=True)
    selected_theme = st.selectbox(
        "Dashboard theme",
        list(THEMES.keys()),
        index=list(THEMES.keys()).index(st.session_state.ui_theme),
        key="theme_picker",
    )
    st.session_state.ui_theme = selected_theme

t = THEMES[st.session_state.ui_theme]

st.markdown(
    f"""
<style>
:root {{
  --tf-bg:{t['bg']}; --tf-panel:{t['panel']}; --tf-panel2:{t['panel2']}; --tf-text:{t['text']};
  --tf-muted:{t['muted']}; --tf-accent:{t['accent']}; --tf-accent2:{t['accent2']};
  --tf-border:{t['border']}; --tf-hero1:{t['hero1']}; --tf-hero2:{t['hero2']}; --tf-glow:{t['glow']};
}}

html, body, [data-testid="stAppViewContainer"], [data-testid="stApp"] {{
  background: var(--tf-bg) !important;
  color: var(--tf-text) !important;
}}
[data-testid="stHeader"] {{background: transparent !important;}}
[data-testid="stToolbar"] {{right: 1rem;}}
[data-testid="stSidebar"] {{
  background: linear-gradient(180deg, var(--tf-panel), var(--tf-bg)) !important;
  border-right: 1px solid var(--tf-border);
}}
[data-testid="stSidebar"] * {{color: var(--tf-text);}}

.block-container {{max-width: 1180px; padding-top: 1.45rem; padding-bottom: 4rem;}}

.hero {{
  position: relative; overflow: hidden; padding: 1.8rem 1.9rem; border: 1px solid var(--tf-border);
  border-radius: 26px; margin-bottom: 1.25rem;
  background: linear-gradient(135deg, var(--tf-hero1), var(--tf-hero2));
  box-shadow: 0 24px 70px var(--tf-glow);
}}
.hero:after {{
  content: ""; position:absolute; width:260px; height:260px; border-radius:50%; right:-95px; top:-125px;
  background: radial-gradient(circle, var(--tf-accent2) 0%, transparent 67%); opacity:.22; filter: blur(8px);
}}
.hero-top {{display:flex; align-items:center; gap:.65rem; margin-bottom:.7rem; position:relative; z-index:1;}}
.brand-dot {{width:10px; height:10px; border-radius:50%; background:var(--tf-accent2); box-shadow:0 0 18px var(--tf-accent2);}}
.brand-chip {{font-size:.75rem; letter-spacing:.12em; text-transform:uppercase; color:var(--tf-muted); font-weight:800;}}
.hero h1 {{margin: 0; font-size: clamp(2rem, 4vw, 3.35rem); line-height:1.03; letter-spacing:-.04em; position:relative; z-index:1;}}
.hero p {{margin: .85rem 0 0 0; color:var(--tf-muted); max-width:780px; font-size:1.02rem; line-height:1.7; position:relative; z-index:1;}}
.hero-badges {{display:flex; flex-wrap:wrap; gap:.55rem; margin-top:1.1rem; position:relative; z-index:1;}}
.hero-badge {{padding:.42rem .72rem; border:1px solid var(--tf-border); border-radius:999px; background:rgba(255,255,255,.04); font-size:.8rem; font-weight:700;}}

.side-kicker {{font-size:.68rem; letter-spacing:.16em; font-weight:800; color:var(--tf-muted); margin:.15rem 0 .4rem;}}
.section-label {{font-size:.72rem; text-transform:uppercase; letter-spacing:.14em; color:var(--tf-muted); font-weight:800; margin-bottom:.25rem;}}
.empty-card {{padding:1.4rem; border:1px dashed var(--tf-border); border-radius:20px; background:var(--tf-panel); color:var(--tf-muted); margin-top:.8rem;}}
.empty-card strong {{color:var(--tf-text);}}

/* Cards / containers */
[data-testid="stMetric"] {{
  background: var(--tf-panel); border:1px solid var(--tf-border); border-radius:18px; padding: .9rem 1rem;
  box-shadow: 0 10px 30px rgba(0,0,0,.10);
}}
[data-testid="stMetricLabel"] {{color:var(--tf-muted) !important;}}
[data-testid="stMetricValue"] {{color:var(--tf-text) !important; font-size:1.15rem !important;}}
[data-testid="stFileUploader"] {{
  background:var(--tf-panel); border:1px solid var(--tf-border); border-radius:20px; padding:.35rem .65rem;
}}
[data-testid="stFileUploaderDropzone"] {{
  border:1px dashed var(--tf-border) !important; background:var(--tf-panel2) !important; border-radius:16px !important;
}}

/* Inputs */
.stTextInput input, .stTextArea textarea, [data-baseweb="select"] > div {{
  background:var(--tf-panel2) !important; color:var(--tf-text) !important; border-color:var(--tf-border) !important;
}}
.stTextArea textarea {{border-radius:14px !important;}}

/* Buttons */
.stButton > button, .stDownloadButton > button {{
  border-radius:13px !important; border:1px solid var(--tf-border) !important;
  background:var(--tf-panel2) !important; color:var(--tf-text) !important; font-weight:750 !important;
  transition:transform .16s ease, border-color .16s ease, box-shadow .16s ease;
}}
.stButton > button:hover, .stDownloadButton > button:hover {{
  transform:translateY(-1px); border-color:var(--tf-accent) !important; box-shadow:0 10px 26px var(--tf-glow);
}}
.stButton > button[kind="primary"] {{
  background:linear-gradient(90deg, var(--tf-accent), var(--tf-accent2)) !important;
  color:white !important; border:0 !important; min-height:3rem;
}}

/* Tabs */
.stTabs [data-baseweb="tab-list"] {{gap:.4rem; border-bottom:1px solid var(--tf-border);}}
.stTabs [data-baseweb="tab"] {{
  height:2.75rem; padding:0 .9rem; border-radius:10px 10px 0 0; color:var(--tf-muted); font-weight:750;
}}
.stTabs [aria-selected="true"] {{color:var(--tf-text) !important; background:var(--tf-panel) !important;}}
.stTabs [data-baseweb="tab-highlight"] {{background:var(--tf-accent) !important;}}

/* Status / alerts */
[data-testid="stAlert"] {{border-radius:14px !important; border:1px solid var(--tf-border) !important;}}
[data-testid="stStatusWidget"] {{border-radius:16px !important; border:1px solid var(--tf-border) !important; background:var(--tf-panel) !important;}}
hr {{border-color:var(--tf-border) !important;}}

.small-muted {{color:var(--tf-muted); font-size:.9rem;}}
.output-card {{border:1px solid var(--tf-border); border-radius:16px; padding:1rem; background:var(--tf-panel);}}

/* Streamlit text colors */
p, label, .stMarkdown, .stCaption, [data-testid="stWidgetLabel"] {{color:var(--tf-text);}}
.stCaption, small {{color:var(--tf-muted) !important;}}

@media (max-width: 800px) {{
  .block-container {{padding-left:1rem; padding-right:1rem;}}
  .hero {{padding:1.35rem; border-radius:20px;}}
  .hero h1 {{font-size:2rem;}}
}}
</style>
""",
    unsafe_allow_html=True,
)

if not render_access_gate():
    st.stop()

st.markdown(
    f"""
<div class="hero">
  <div class="hero-top"><span class="brand-dot"></span><span class="brand-chip">Team Fahad AI Studio</span></div>
  <h1>Turn media into usable content.</h1>
  <p>Upload once. Get a polished transcript, speaker-aware subtitles, summaries, translations and creator-ready deliverables from one clean workspace.</p>
  <div class="hero-badges">
    <span class="hero-badge">Audio + Video</span>
    <span class="hero-badge">Speaker Detection</span>
    <span class="hero-badge">SRT / VTT</span>
    <span class="hero-badge">AI Summary</span>
    <span class="hero-badge">Translation</span>
  </div>
</div>
""",
    unsafe_allow_html=True,
)

api_key = secret_or_env("GEMINI_API_KEY")
if not api_key:
    st.warning("Developer setup: no GEMINI_API_KEY was found in Streamlit Secrets or environment variables.")
    api_key = st.text_input("Gemini API key for this local session", type="password")

with st.sidebar:
    st.markdown("<div class='side-kicker' style='margin-top:1.2rem'>TRANSCRIPTION</div>", unsafe_allow_html=True)
    st.subheader("Settings")
    language_name = st.selectbox("Spoken language", list(LANGUAGES.keys()), index=0)
    mode = st.radio(
        "Output mode",
        ["Smart clean transcript", "Detailed subtitles + speakers", "Exact verbatim transcript"],
        help=(
            "Smart cleans filler words and formatting. Detailed enables speaker labels and word timestamps. "
            "Exact verbatim preserves what was spoken."
        ),
    )

    custom_vocab_text = ""
    if mode != "Detailed subtitles + speakers":
        custom_vocab_text = st.text_area(
            "Custom vocabulary (optional)",
            placeholder="Tamrin Institute, Team Fahad, product names...",
            help="One term per line or comma-separated. Best for uncommon names and technical words.",
        )
    else:
        st.caption("Custom vocabulary is disabled in detailed speaker/timestamp mode.")

    st.divider()
    st.caption("Privacy: temporary local files are deleted after processing, and the Gemini Files API copy is deleted after the request.")

st.markdown("<div class='section-label'>01 · Upload & process</div>", unsafe_allow_html=True)
st.markdown("### Start with your media file")
st.caption("Drop an audio or video file below. Your selected transcription mode and language settings are applied automatically.")

uploaded = st.file_uploader(
    "Upload audio or video",
    type=["mp3", "wav", "m4a", "aac", "ogg", "flac", "mp4", "mov", "mkv"],
    help=f"Maximum app file size: {MAX_UPLOAD_MB} MB.",
)

if uploaded is None:
    st.markdown(
        """
        <div class="empty-card">
          <strong>Ready when you are.</strong><br>
          Choose a file to unlock transcription, speaker subtitles, AI summaries, translation and client-ready downloads.
        </div>
        """,
        unsafe_allow_html=True,
    )

if uploaded is not None:
    size_mb = getattr(uploaded, "size", 0) / (1024 * 1024)
    c1, c2, c3 = st.columns(3)
    c1.metric("File", uploaded.name)
    c2.metric("Size", f"{size_mb:.1f} MB")
    c3.metric("Mode", "Detailed" if mode.startswith("Detailed") else "Standard")

    vocab = [item.strip() for item in re.split(r"[,\n]", custom_vocab_text) if item.strip()]
    if len(vocab) > 100:
        st.info("Only the first 100 custom vocabulary terms will be sent for best results.")

    if mode == "Detailed subtitles + speakers":
        st.info("Detailed mode is intended for recordings up to about 30 minutes. Standard modes support up to about 60 minutes per request.")

    if st.button("Transcribe file", type="primary", disabled=not bool(api_key), use_container_width=True):
        reset_outputs()
        if size_mb > MAX_UPLOAD_MB:
            st.error(f"Please upload a file smaller than {MAX_UPLOAD_MB} MB.")
        else:
            status = st.status("Preparing media…", expanded=True)
            try:
                status.write("Extracting/reading audio…")
                status.write("Uploading securely for transcription…")
                status.write("Running speech-to-text…")
                result = transcribe_media(
                    uploaded=uploaded,
                    api_key=api_key,
                    language_codes=LANGUAGES[language_name],
                    mode=mode,
                    custom_vocabulary=vocab,
                )
                st.session_state.result = result
                st.session_state.source_name = uploaded.name
                status.update(label="Transcription complete", state="complete", expanded=False)
            except errors.APIError as exc:
                status.update(label="Transcription failed", state="error")
                code = getattr(exc, "code", "API")
                message = getattr(exc, "message", str(exc))
                st.error(f"Gemini request failed ({code}): {message}")
            except Exception as exc:
                status.update(label="Transcription failed", state="error")
                st.error(str(exc))

result = st.session_state.get("result")
if result:
    source_name = st.session_state.get("source_name", "transcript")
    base_name = safe_name(source_name)

    if result.get("duration"):
        st.success(f"Completed • Media duration: {format_clock(result['duration'])}")
    else:
        st.success("Completed")

    st.markdown("<div class='section-label' style='margin-top:1.2rem'>02 · Workspace</div>", unsafe_allow_html=True)
    st.markdown("### Your transcription workspace")

    tab_transcript, tab_subtitles, tab_ai, tab_downloads = st.tabs(
        ["Transcript", "Speakers & subtitles", "AI client tools", "Downloads"]
    )

    with tab_transcript:
        st.subheader("Transcript")
        st.text_area("Transcript text", result.get("transcript", ""), height=430)
        st.download_button(
            "Download TXT",
            result.get("transcript", ""),
            file_name=f"{base_name}_transcript.txt",
            mime="text/plain",
        )

    with tab_subtitles:
        if result.get("segments"):
            st.subheader("Speaker transcript")
            st.text_area("Speaker + timestamp transcript", result.get("speaker_transcript", ""), height=330)
            d1, d2, d3 = st.columns(3)
            d1.download_button(
                "Download speaker TXT",
                result.get("speaker_transcript", ""),
                file_name=f"{base_name}_speaker_transcript.txt",
                mime="text/plain",
                use_container_width=True,
            )
            d2.download_button(
                "Download SRT",
                result.get("srt", ""),
                file_name=f"{base_name}.srt",
                mime="application/x-subrip",
                use_container_width=True,
            )
            d3.download_button(
                "Download VTT",
                result.get("vtt", ""),
                file_name=f"{base_name}.vtt",
                mime="text/vtt",
                use_container_width=True,
            )
            st.caption(f"Subtitle segments: {len(result.get('segments', []))}")
        else:
            st.info("Speaker labels and subtitle files are created when you use “Detailed subtitles + speakers” mode.")

    with tab_ai:
        st.subheader("Turn the transcript into client-ready deliverables")
        st.caption("These are generated only when you click a button, so you control extra API usage.")

        col_a, col_b = st.columns(2)
        with col_a:
            if st.button("Generate summary + key points", use_container_width=True, disabled=not bool(api_key)):
                with st.spinner("Creating summary…"):
                    try:
                        st.session_state.summary = generate_summary(api_key, result)
                    except Exception as exc:
                        st.error(f"Summary failed: {exc}")
        with col_b:
            if st.button("Generate creator content pack", use_container_width=True, disabled=not bool(api_key)):
                with st.spinner("Creating content pack…"):
                    try:
                        st.session_state.content_pack = generate_content_pack(api_key, result)
                    except Exception as exc:
                        st.error(f"Content pack failed: {exc}")

        target = st.selectbox("Translation target", TRANSLATION_LANGUAGES)
        if st.button(f"Translate full transcript to {target}", use_container_width=True, disabled=not bool(api_key)):
            with st.spinner(f"Translating to {target}…"):
                try:
                    st.session_state.translation = generate_translation(api_key, result, target)
                    st.session_state.translation_target = target
                except Exception as exc:
                    st.error(f"Translation failed: {exc}")

        if st.session_state.get("summary"):
            st.markdown("### Summary & key points")
            st.text_area("Summary", st.session_state.summary, height=300)
        if st.session_state.get("translation"):
            label = st.session_state.get("translation_target", "Translation")
            st.markdown(f"### {label} translation")
            st.text_area("Translated transcript", st.session_state.translation, height=360)
        if st.session_state.get("content_pack"):
            st.markdown("### Creator content pack")
            st.text_area("Content pack", st.session_state.content_pack, height=430)

    with tab_downloads:
        extras = {
            "summary": st.session_state.get("summary", ""),
            f"translation_{str(st.session_state.get('translation_target', '')).lower()}": st.session_state.get("translation", ""),
            "content_pack": st.session_state.get("content_pack", ""),
        }
        sections = [("Transcript", result.get("transcript", ""))]
        if result.get("speaker_transcript"):
            sections.append(("Speaker Transcript", result.get("speaker_transcript", "")))
        if st.session_state.get("summary"):
            sections.append(("Summary & Key Points", st.session_state.summary))
        if st.session_state.get("translation"):
            sections.append((f"Translation - {st.session_state.get('translation_target', '')}", st.session_state.translation))
        if st.session_state.get("content_pack"):
            sections.append(("Creator Content Pack", st.session_state.content_pack))

        docx_data = make_docx_bytes(APP_TITLE, sections)
        zip_data = make_zip_bytes(base_name, result, extras)

        st.write("Download a polished document or one ZIP containing every output currently generated.")
        dl1, dl2 = st.columns(2)
        dl1.download_button(
            "Download complete DOCX",
            docx_data,
            file_name=f"{base_name}_complete.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )
        dl2.download_button(
            "Download everything as ZIP",
            zip_data,
            file_name=f"{base_name}_team_fahad_pack.zip",
            mime="application/zip",
            use_container_width=True,
        )

st.divider()
st.markdown(
    "<div class='small-muted' style='text-align:center;padding:.4rem 0 1rem'>Team Fahad AI Studio · Client-ready transcription workflow</div>",
    unsafe_allow_html=True,
)